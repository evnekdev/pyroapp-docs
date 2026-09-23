from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference, Series
from openpyxl.workbook.properties import CalcProperties
import json, re

ROOT = Path(__file__).resolve().parents[1] if 'scripts' in Path(__file__).parts else Path('/mnt/data/pyroapp_examples')
OUT = ROOT / 'docs' / 'tutorial-assets' / 'workbooks' if (ROOT / 'docs').exists() else Path('/mnt/data/pyroapp_examples/compact_out')
OUT.mkdir(parents=True, exist_ok=True)
DOC='https://evnekdev.github.io/pyroapp-docs/'
DATA=DOC+'tutorial-assets/datafiles/'
NAVY='1F4E78'; WHITE='FFFFFF'; BLUE='D9EAF7'; GRAY='E7E6E6'; ORANGE='FCE4D6'; RED='F4CCCC'; PURPLE='E4DFEC'; TEAL='DDEBF7'
CELL_EDGE = Side(style='thin', color='B7C9D6')
CELL_BORDER = Border(left=CELL_EDGE, right=CELL_EDGE, top=CELL_EDGE, bottom=CELL_EDGE)

def sh(ws, widths=None):
    # Filled cells do not display Excel gridlines.  Keep gridlines enabled and
    # add explicit borders to populated cells so the teaching layout never
    # hides a cell boundary.
    ws.sheet_view.showGridLines=True
    for c,w in (widths or {'A':24,'B':20,'C':18,'D':18,'E':18,'F':18,'G':18,'H':18,'I':18,'J':18,'K':18,'L':18,'M':18,'N':18,'O':18,'P':18}).items(): ws.column_dimensions[c].width=w

def head(ws,text,sub=''):
    ws['A1']=text; ws['A1'].font=Font(size=18,bold=True,color=WHITE); ws['A1'].fill=PatternFill('solid',fgColor=NAVY); ws.merge_cells('A1:J1')
    if sub: ws['A2']=sub; ws['A2'].alignment=Alignment(wrap_text=True,vertical='top'); ws.merge_cells('A2:J2'); ws.row_dimensions[2].height=42

def section(ws,row,text,end=8):
    ws.cell(row,1,text); ws.cell(row,1).font=Font(bold=True,color=WHITE); ws.cell(row,1).fill=PatternFill('solid',fgColor=NAVY); ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=end)

def inp(c): c.font=Font(color='0000FF'); c.fill=PatternFill('solid',fgColor=BLUE)
def ctl(c): c.font=Font(color='7030A0'); c.fill=PatternFill('solid',fgColor=PURPLE)
def th(c): c.font=Font(bold=True,color=WHITE); c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(horizontal='center',wrap_text=True)
def note(ws,row,text,end=8,color=TEAL):
    ws.cell(row,1,text); ws.cell(row,1).fill=PatternFill('solid',fgColor=color); ws.cell(row,1).alignment=Alignment(wrap_text=True,vertical='top'); ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=end)
    ws.row_dimensions[row].height=max(30, 15*(1+len(text)//90))
def hlink(c,label,url): c.value=label; c.hyperlink=url; c.font=Font(color='008000',underline='single')

def calc(wb): wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True,calcId=191029)

def start(wb,title,datafile,url,desc):
    ws=wb.create_sheet('Start'); sh(ws,{'A':26,'B':50,'C':18,'D':18,'E':18,'F':18,'G':18,'H':18,'I':18,'J':18}); head(ws,title,desc)
    section(ws,4,'Before you calculate')
    steps=[f'1. Install/load PyroApp in desktop Excel.',f'2. Download {datafile} and place it beside this workbook.','3. Blue cells are inputs; black cells are formulas; purple cells are controls.','4. The workbooks use literal text "" in unused CA_CALCULATE header cells (no leading equals sign).']
    for r,t in enumerate(steps,5): ws.cell(r,1,t); ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=8)
    ws['A10']='Datafile'; ws['B10']=datafile; inp(ws['B10']); ws['A11']='Download DAT'; hlink(ws['B11'],datafile,url); ws['A12']='Docs'; hlink(ws['B12'],'PyroApp documentation',DOC); ws['A14']='PyroApp status'; ws['B14']='=IFERROR(XLL_CA_VERSION(),"PyroApp not loaded")'
    return ws

def hdr(ws,row,col,defs):
    for j,(code,phase,comp) in enumerate(defs,col):
        ws.cell(row,j,code); th(ws.cell(row,j)); ws.cell(row+1,j,phase if phase is not None else '""'); ws.cell(row+2,j,comp if comp is not None else '""')
        ws.cell(row+1,j).fill=PatternFill('solid',fgColor=GRAY); ws.cell(row+2,j).fill=PatternFill('solid',fgColor=GRAY)

def chart3(ws,anchor,cats_col,cols,start,end,title,xlab):
    ch=LineChart(); ch.title=title; ch.y_axis.title='Amount'; ch.x_axis.title=xlab; cats=Reference(ws,min_col=cats_col,min_row=start,max_row=end)
    for col,name in cols:
        ch.append(Series(Reference(ws,min_col=col,min_row=start,max_row=end),title=name))
    ch.set_categories(cats); ch.height=8; ch.width=14; ws.add_chart(ch,anchor)

def getting_started():
    wb=Workbook(); wb.remove(wb.active); calc(wb); start(wb,'PyroApp Example 01 — Getting Started (Ca-Zn-O)','Ca-Zn-O.dat',DATA+'examples/Ca-Zn-O.dat','Runtime checks, DAT discovery, basis conversion, mesh generation, equilibrium tables, sweeps and phase targets.')
    ws=wb.create_sheet('Runtime'); sh(ws); head(ws,'Runtime and licence information');
    for r,(a,b) in enumerate([('Version','=XLL_CA_VERSION()'),('Lite?','=XLL_CA_ISLITE()'),('User ID','=XLL_CA_USER_ID()'),('Licence holder','=XLL_CA_LICENSE_HOLDER_NAME()'),('Program ID','=XLL_CA_PROGRAM_ID()'),('Maximum dimensions','=XLL_CA_DIMENSIONS_MAX()'),('DLLs','=XLL_CA_CHEMAPPDLLS()'),('DAT dimensions','=XLL_CA_DIMENSIONS(Start!$B$10)')],5): ws.cell(r,1,a); ws.cell(r,2,b)
    ws=wb.create_sheet('Explore'); sh(ws,{'A':22,'B':26,'C':4,'D':28,'E':4,'F':28,'G':4,'H':28,'I':4,'J':34,'K':4,'L':34}); head(ws,'Explore the datafile','LIST functions return exact DAT identities.')
    for hc,label,fc,form in [('A4','Components','A5','=XLL_CA_LIST_COMPONENTS(Start!$B$10)'),('C4','Phases','C5','=XLL_CA_LIST_PHASES(Start!$B$10)'),('E4','Solutions','E5','=XLL_CA_LIST_SOLUTIONS(Start!$B$10)'),('G4','Compounds','G5','=XLL_CA_LIST_COMPOUNDS(Start!$B$10)'),('I4','Slag-liq constituents','I5','=XLL_CA_LIST_CONSTITUENTS(Start!$B$10,"Slag-liq")'),('K4','Slag-liq species','K5','=XLL_CA_LIST_SPECIES(Start!$B$10,"Slag-liq")')]: ws[hc]=label; th(ws[hc]); ws[fc]=form
    section(ws,26,'Direct property lookups',6); ws['A28']='Zn molar mass'; ws['B28']='=XLL_CA_GET_COMPONENT_WEIGHTS(Start!$B$10,"Zn")'; ws['A29']='Ca molar mass'; ws['B29']='=XLL_CA_GET_COMPONENT_WEIGHTS(Start!$B$10,"Ca")'; ws['A30']='ZnO(s) H298'; ws['B30']='=XLL_CA_GET_COMPOUND_H298(Start!$B$10,"ZnO(s)")'
    ws=wb.create_sheet('Basis + utilities'); sh(ws); head(ws,'Composition-basis lab and mesh generation','Each block has one clearly labelled conversion. Blue cells are inputs; the selected formula cell spills the converted composition.')
    section(ws,4,'1. One molar-fraction row: CaO + ZnO → elemental Ca, Zn, O',7)
    for c,v in [('B5','CaO'),('C5','ZnO'),('E5','Ca'),('F5','Zn'),('G5','O')]: ws[c]=v; th(ws[c])
    ws['A7']='Input mole fractions'; ws['B7']=0.40; ws['C7']=0.60; inp(ws['B7']); inp(ws['C7']); ws['D7']='Normalized elemental mole fractions'; ws['E7']='=XLL_DATA_CHANGE_BASIS($B$5:$C$5,$E$5:$G$5,$B$7:$C$7,FALSE,FALSE,TRUE)'
    section(ws,11,'2. Multiple molar-fraction rows: CaO + ZnO → elemental fractions',7)
    for r,x in enumerate([i/10 for i in range(11)],14): ws.cell(r,2,x); ws.cell(r,3,1-x); inp(ws.cell(r,2)); inp(ws.cell(r,3))
    ws['A13']='Input row'; ws['B13']='CaO'; ws['C13']='ZnO'; ws['D13']='Ca'; ws['E13']='Zn'; ws['F13']='O'
    for c in ['A13','B13','C13','D13','E13','F13']: th(ws[c])
    ws['D14']='=XLL_DATA_CHANGE_BASIS($B$5:$C$5,$E$5:$G$5,$B$14:$C$24,FALSE,FALSE,TRUE)'
    section(ws,27,'3. Molar amounts: CaO + ZnO → elemental Ca, Zn, O amounts',7)
    for c,v in [('B28','CaO'),('C28','ZnO'),('E28','Ca'),('F28','Zn'),('G28','O')]: ws[c]=v; th(ws[c])
    ws['A30']='Formula-unit amounts (mol)'; ws['B30']=2; ws['C30']=1; inp(ws['B30']); inp(ws['C30']); ws['D30']='Element amounts (mol)'; ws['E30']='=XLL_DATA_CHANGE_BASIS($B$28:$C$28,$E$28:$G$28,$B$30:$C$30,FALSE,FALSE,FALSE)'
    section(ws,34,'4. Mass amounts: CaO + ZnO → elemental mole amounts (not normalized)',7)
    for c,v in [('B35','CaO'),('C35','ZnO'),('E35','Ca'),('F35','Zn'),('G35','O')]: ws[c]=v; th(ws[c])
    ws['A37']='Mass amounts (g)'; ws['B37']=112.16; ws['C37']=81.38; inp(ws['B37']); inp(ws['C37']); ws['D37']='Element amounts (mol)'; ws['E37']='=XLL_DATA_CHANGE_BASIS($B$35:$C$35,$E$35:$G$35,$B$37:$C$37,TRUE,FALSE,FALSE)'
    section(ws,41,'5. Molar amounts: CaO + ZnO → normalized mass fractions',7)
    for c,v in [('B42','CaO'),('C42','ZnO'),('E42','Ca'),('F42','Zn'),('G42','O')]: ws[c]=v; th(ws[c])
    ws['A44']='Formula-unit amounts (mol)'; ws['B44']=2; ws['C44']=1; inp(ws['B44']); inp(ws['C44']); ws['D44']='Element mass fractions'; ws['E44']='=XLL_DATA_CHANGE_BASIS($B$42:$C$42,$E$42:$G$42,$B$44:$C$44,FALSE,TRUE,TRUE)'
    ws['I4']='Uniform simplex mesh'; ws.merge_cells('I4:L4'); th(ws['I4']);
    for r,(a,b) in enumerate([('nintervals',5),('nactive',2),('ntotal',3)],5): ws.cell(r,9,a); ws.cell(r,10,b); inp(ws.cell(r,10))
    ws['I9']='=XLL_DATA_GENERATE_MESH($J$5,$J$6,$J$7)'
    ws.freeze_panes='A4'
    ws=wb.create_sheet('First equilibrium'); sh(ws); head(ws,'First equilibrium calculation','One 50:50 CaO-ZnO composition at 1200 °C and 1 bar. The grey header rows use literal text "" where a phase or component is not needed.')
    hdr(ws,4,2,[('T, [C]',None,None),('P, [bar]',None,None),('IA',None,'Ca'),('IA',None,'Zn'),('IA',None,'O')])
    for j,v in enumerate([1200,1,0.5,0.5,1.0],2): ws.cell(7,j,v); inp(ws.cell(7,j))
    hdr(ws,4,8,[('A','Slag-liq',None),('A','Monoxide',None),('A','Zincite',None),('ERROR',None,None),('NSTABLE',None,None)])
    ws['H7']='=XLL_CA_CALCULATE(Start!$B$10,$B$4:$F$6,$B$7:$F$7,$H$4:$L$6)'; note(ws,10,'Select H7 to read the complete formula in Excel’s Formula Bar. A zero ERROR result means the row solved successfully.',12,TEAL)
    ws=wb.create_sheet('Equilibrium'); sh(ws); head(ws,'Equilibrium table','Eleven independent CaO-ZnO compositions at 1200 °C and 1 bar.')
    hdr(ws,4,2,[('T, [C]',None,None),('P, [bar]',None,None),('IA',None,'Ca'),('IA',None,'Zn'),('IA',None,'O')])
    for r,x in enumerate([i/10 for i in range(11)],7):
        for j,v in enumerate([1200,1,x,1-x,1.0],2): ws.cell(r,j,v); inp(ws.cell(r,j))
    hdr(ws,4,8,[('T, [C]',None,None),('A','Slag-liq',None),('A','Monoxide',None),('A','Zincite',None),('XP','Slag-liq','Zn'),('ERROR',None,None),('NSTABLE',None,None)]); ws['H7']='=XLL_CA_CALCULATE(Start!$B$10,$B$4:$F$6,$B$7:$F$17,$H$4:$N$6)'; chart3(ws,'P4',4,[(9,'Slag-liq'),(10,'Monoxide'),(11,'Zincite')],7,17,'Phase amounts at 1200 °C','CaO fraction')
    ws=wb.create_sheet('Temperature sweep'); sh(ws); head(ws,'Temperature sweep','50:50 CaO-ZnO from 700 to 1800 °C.')
    hdr(ws,4,2,[('T, [C]',None,None),('P, [bar]',None,None),('IA',None,'Ca'),('IA',None,'Zn'),('IA',None,'O')]); temps=list(range(700,1801,100))
    for r,T in enumerate(temps,7):
        for j,v in enumerate([T,1,0.5,0.5,1.0],2): ws.cell(r,j,v); inp(ws.cell(r,j))
    hdr(ws,4,8,[('T, [C]',None,None),('A','Slag-liq',None),('A','Monoxide',None),('A','Zincite',None),('ERROR',None,None),('NSTABLE',None,None)]); ws['H7']=f'=XLL_CA_CALCULATE(Start!$B$10,$B$4:$F$6,$B$7:$F${6+len(temps)},$H$4:$M$6)'; chart3(ws,'O4',8,[(9,'Slag-liq'),(10,'Monoxide'),(11,'Zincite')],7,6+len(temps),'Phase amounts vs temperature','Temperature [°C]')
    ws=wb.create_sheet('Phase selection'); sh(ws); head(ws,'Phase selection','The optional entered-phase range makes the selection explicit before the calculation is solved.')
    hdr(ws,4,2,[('T, [C]',None,None),('P, [bar]',None,None),('IA',None,'Ca'),('IA',None,'Zn'),('IA',None,'O')])
    for j,v in enumerate([1200,1,0.5,0.5,1.0],2): ws.cell(7,j,v); inp(ws.cell(7,j))
    hdr(ws,4,8,[('A','Slag-liq',None),('A','Monoxide',None),('A','Zincite',None),('ERROR',None,None),('NSTABLE',None,None)])
    ws['N4']='Entered phases'; th(ws['N4']); ws['N5']='Slag-liq'; ws['N6']='Monoxide'; inp(ws['N5']); inp(ws['N6']); ws['H7']='=XLL_CA_CALCULATE(Start!$B$10,$B$4:$F$6,$B$7:$F$7,$H$4:$L$6,$N$5:$N$6)'; note(ws,10,'The entered range is intentionally visible. Change its phase names only when the physical model justifies restricting the phase set.',14,ORANGE)
    ws=wb.create_sheet('Phase target'); sh(ws); head(ws,'Formation-temperature target','Solve T for Monoxide formation rather than scanning temperature.')
    hdr(ws,4,2,[('P, [bar]',None,None),('IA',None,'Ca'),('IA',None,'Zn'),('IA',None,'O'),('FORMATION',None,None),('TLOW',None,None),('THIGH',None,None)])
    for r,x in enumerate([0.1,0.3,0.5,0.7,0.9],7):
        for j,v in enumerate([1,x,1-x,1.0,'Monoxide',500,2200],2): ws.cell(r,j,v); inp(ws.cell(r,j))
    hdr(ws,4,11,[('T, [C]',None,None),('ERROR',None,None),('NSTABLE',None,None),('AC','Monoxide',None)]); ws['P4']='Entered phases'; th(ws['P4']); ws['P5']='Slag-liq'; ws['P6']='Monoxide'; inp(ws['P5']); inp(ws['P6']); ws['K7']='=XLL_CA_CALCULATE(Start!$B$10,$B$4:$H$6,$B$7:$H$11,$K$4:$N$6,$P$5:$P$6)'; note(ws,14,'This teaching target deliberately restricts the entered phases to Slag-liq + Monoxide. Real phase selection must match the physical problem.',14,ORANGE)
    return save(wb,'01_PyroApp_Getting_Started_Ca-Zn-O.xlsx')

def dat_parameters():
    wb=Workbook(); wb.remove(wb.active); calc(wb); start(wb,'PyroApp Example 02 — DAT Inspection and Editing (Ca-Zn-O)','Ca-Zn-O.dat',DATA+'examples/Ca-Zn-O.dat','Browse an open DAT, read compound/constituent parameters, inspect interactions and use a gated SET operation.')
    ws=wb.create_sheet('Browse'); sh(ws); head(ws,'Browse an open DAT');
    for hc,label,fc,form in [('A4','Components','A5','=XLL_CA_LIST_COMPONENTS(Start!$B$10)'),('C4','Solutions','C5','=XLL_CA_LIST_SOLUTIONS(Start!$B$10)'),('E4','Compounds','E5','=XLL_CA_LIST_COMPOUNDS(Start!$B$10)'),('G4','Phases','G5','=XLL_CA_LIST_PHASES(Start!$B$10)')]: ws[hc]=label; th(ws[hc]); ws[fc]=form
    section(ws,30,'Inspect Slag-liq',8); ws['A32']='Constituents'; th(ws['A32']); ws['B32']='=XLL_CA_LIST_CONSTITUENTS(Start!$B$10,"Slag-liq")'; ws['D32']='Species'; th(ws['D32']); ws['E32']='=XLL_CA_LIST_SPECIES(Start!$B$10,"Slag-liq")'; ws['G32']='G interactions'; th(ws['G32']); ws['H32']='=XLL_CA_LIST_INTERACTIONS_G(Start!$B$10,"Slag-liq")'
    ws=wb.create_sheet('Compound data'); sh(ws); head(ws,'Compound parameter lookups'); ws['A4']='Phase'; ws['B4']='CaO(s)'; inp(ws['B4'])
    for r,(a,b) in enumerate([('H298','=XLL_CA_GET_COMPOUND_H298(Start!$B$10,$B$4)'),('S298','=XLL_CA_GET_COMPOUND_S298(Start!$B$10,$B$4)'),('Molar mass','=XLL_CA_GET_COMPOUND_WEIGHTS(Start!$B$10,$B$4)'),('Stoichiometry','=XLL_CA_GET_COMPOUND_STOICHIOMETRY(Start!$B$10,$B$4)'),('Range count','=XLL_CA_GET_COMPOUND_RANGE_COUNT(Start!$B$10,$B$4)'),('Range 1 Tupper','=XLL_CA_GET_COMPOUND_TUPPER(Start!$B$10,$B$4,1)'),('Range 1 coefficient 1','=XLL_CA_GET_COMPOUND_CP(Start!$B$10,$B$4,1,1)')],6): ws.cell(r,1,a); ws.cell(r,2,b)
    ws=wb.create_sheet('Constituent data'); sh(ws); head(ws,'Solution-constituent parameter lookups'); ws['A4']='Phase'; ws['B4']='Slag-liq'; inp(ws['B4']); ws['A5']='Constituent'; ws['B5']='CaO'; inp(ws['B5'])
    for r,(a,b) in enumerate([('H298','=XLL_CA_GET_CONSTITUENT_H298(Start!$B$10,$B$4,$B$5)'),('S298','=XLL_CA_GET_CONSTITUENT_S298(Start!$B$10,$B$4,$B$5)'),('Molar mass','=XLL_CA_GET_CONSTITUENT_WEIGHTS(Start!$B$10,$B$4,$B$5)'),('Stoichiometry','=XLL_CA_GET_CONSTITUENT_STOICHIOMETRY(Start!$B$10,$B$4,$B$5)'),('Range count','=XLL_CA_GET_CONSTITUENT_RANGE_COUNT(Start!$B$10,$B$4,$B$5)'),('Range 1 Tupper','=XLL_CA_GET_CONSTITUENT_TUPPER(Start!$B$10,$B$4,$B$5,1)'),('Range 1 coefficient 1','=XLL_CA_GET_CONSTITUENT_CP(Start!$B$10,$B$4,$B$5,1,1)')],7): ws.cell(r,1,a); ws.cell(r,2,b)
    ws=wb.create_sheet('Interactions'); sh(ws,{'A':24,'B':54,'C':18,'D':18,'E':18,'F':18,'G':18}); head(ws,'Excess Gibbs-energy interactions'); ws['A4']='Phase'; ws['B4']='Slag-liq'; inp(ws['B4']); ws['A6']='Interaction list'; th(ws['A6']); ws['B6']='=XLL_CA_LIST_INTERACTIONS_G(Start!$B$10,$B$4)'; section(ws,22,'Read the first interaction',7); ws['A24']='Interaction'; ws['B24']='=B6'; ws['A25']='Index'; ws['B25']='=XLL_CA_GET_INTERACTION_INDICES(Start!$B$10,$B$4,$B$24)'
    for j,h in enumerate(['constant','T','T ln(T)','T^2','T^3','1/T'],2): ws.cell(28,j,h); th(ws.cell(28,j))
    ws['B29']='=XLL_CA_GET_INTERACTION_PARAMETERS_G(Start!$B$10,$B$4,$B$24,0)'
    ws=wb.create_sheet('Safe edit'); sh(ws); head(ws,'Gated SET example'); note(ws,4,'SET modifies the DAT file. Leave Enable SET = FALSE until you intentionally want to write the Proposed value. Keep an untouched reference DAT.',6,RED)
    vals=[('Phase','Slag-liq'),('Interaction','=Interactions!B24'),('Value index',1),('Current value','=XLL_CA_GET_INTERACTION_PARAMETERS_G(Start!$B$10,$B$8,$B$9,$B$10,$B$13)'),('Proposed value','=B11'),('Update token',0),('Enable SET',False)]
    for r,(a,b) in enumerate(vals,8): ws.cell(r,1,a); ws.cell(r,2,b)
    inp(ws['B8']); inp(ws['B10']); inp(ws['B12']); ctl(ws['B13']); ctl(ws['B14']); ws['A16']='SET result'; ws['B16']='=IF($B$14,XLL_CA_SET_INTERACTION_PARAMETERS_G(Start!$B$10,$B$8,$B$9,$B$10,$B$12,$B$13),"Disabled")'; note(ws,19,'Typical use: read Current value, replace Proposed value with a number, enable SET, increment Update token, then disable SET and recalculate dependent calculations.',6)
    return save(wb,'02_PyroApp_DAT_Inspection_and_Editing_Ca-Zn-O.xlsx')

def optimization():
    wb=Workbook(); wb.remove(wb.active); calc(wb); start(wb,'PyroApp Example 03 — Optimization Workflows (Ca-Zn-O)','Ca-Zn-O.dat',DATA+'examples/Ca-Zn-O.dat','Synthetic regression helpers plus the live derivative-matrix dependency pattern for an open DAT.')
    ws=wb.create_sheet('Synthetic regression'); sh(ws); head(ws,'Synthetic linear-regression example','Illustrative numbers isolate the optimization mathematics; they are not experimental thermodynamic data.');
    for j,h in enumerate(['Category','Residual','Weight','Use?','dT/dx'],1): ws.cell(5,j,h); th(ws.cell(5,j))
    cats=['liquidus','liquidus','invariant','invariant','mixing','mixing']; res=[25,-15,8,-6,120,-80]; w=[1,1,2,2,.05,.05]
    for i in range(6):
        for j,v in enumerate([cats[i],res[i],w[i],True,1],1): ws.cell(6+i,j,v); inp(ws.cell(6+i,j))
    section(ws,14,'Derivative matrix A = d(current)/d(parameter)',6)
    for j,h in enumerate(['p1','p2','p3'],2): ws.cell(15,j,h); th(ws.cell(15,j))
    A=[[.8,.1,0],[1,-.1,.1],[.2,.6,-.1],[.1,.8,.2],[3,.5,1.4],[-2,.4,1]]
    for i,row in enumerate(A,16):
        for j,v in enumerate(row,2): ws.cell(i,j,v); inp(ws.cell(i,j))
    ws['B23']='Free mask'; ws['C23']=True; ws['D23']=True; ws['E23']=False; ws['B24']='Fixed values'; ws['C24']=0; ws['D24']=0; ws['E24']=0
    for c in ['C23','D23','E23','C24','D24','E24']: inp(ws[c])
    ws['B26']='Proposed Δp'; ws['C26']='=XLL_LINEAR_REGRESSION($B$16:$D$21,$B$6:$B$11,$C$6:$C$11,$C$23:$E$23,$C$24:$E$24)'; ws['B28']='Best 2-parameter mask'; ws['C28']='=XLL_BEST_COMBINATION($B$16:$D$21,$B$6:$B$11,$C$6:$C$11,$C$23:$E$23,$C$24:$E$24,2)'
    for r in range(33,39): ws.cell(r,1,f'=B{r-27}-SUMPRODUCT(B{r-17}:D{r-17},$C$26:$E$26)')
    ws['C33']='=XLL_PD_ERROR_TABLE($A$6:$A$11,$B$6:$B$11,$A$33:$A$38,$D$6:$D$11,$E$6:$E$11)'; note(ws,41,'Convention: residual = target - current; A = d(current)/d(parameter); solve A Δp ≈ residual. Recalculate the real nonlinear model after applying a step.',10)
    ws=wb.create_sheet('Derivative matrix setup'); sh(ws); ws['D1']='Live derivative-matrix setup'; ws.merge_cells('D1:J1'); th(ws['D1']); ws['D2']='Run PyroApp → Calculation → Calculate derivative matrix while this sheet is active. B1:B6 are reserved for the command.'; ws.merge_cells('D2:J3'); ws['D2'].alignment=Alignment(wrap_text=True)
    cfg=[('Parameter range','Parameters!$B$8:$B$9'),('Residual/output range','Live targets!$J$8:$J$12'),('Trigger cell','Parameters!$B$12'),('Derivative output top-left','$D$8'),('Step(s)','Parameters!$B$10:$B$11'),('Parameter mask','Parameters!$B$13:$B$14')]
    for r,(a,b) in enumerate(cfg,1): ws.cell(r,1,a); ws.cell(r,2,b); ctl(ws.cell(r,2))
    ws['C7']='Target'; th(ws['C7']); ws['D7']='Parameter 1'; th(ws['D7']); ws['E7']='Parameter 2'; th(ws['E7']);
    for r in range(8,13): ws.cell(r,3,f'Target {r-7}')
    note(ws,15,'PyroApp perturbs numeric parameter cells, changes the trigger, waits for dependent SET/calculation formulas, writes derivative columns, then restores the original parameter vector.',10)
    ws=wb.create_sheet('Parameters'); sh(ws); head(ws,'Optimization parameter block','The starting numbers are the first Slag-liq ordinary G interaction terms in the supplied working DAT.'); ws['A4']='Phase'; ws['B4']='Slag-liq'; inp(ws['B4']); ws['A5']='Interaction identity'; ws['B5']='=XLL_CA_LIST_INTERACTIONS_G(Start!$B$10,$B$4)'; ws['A8']='Parameter 1 constant'; ws['B8']=-119160.32; inp(ws['B8']); ws['A9']='Parameter 2 T'; ws['B9']=21.673120; inp(ws['B9']); ws['A10']='Step p1'; ws['B10']=100; inp(ws['B10']); ws['A11']='Step p2'; ws['B11']=.1; inp(ws['B11']); ws['A12']='Trigger'; ws['B12']=0; ctl(ws['B12']); ws['A13']='Use p1?'; ws['B13']=True; ctl(ws['B13']); ws['A14']='Use p2?'; ws['B14']=True; ctl(ws['B14']); ws['A18']='Write p1'; ws['B18']='=XLL_CA_SET_INTERACTION_PARAMETERS_G(Start!$B$10,$B$4,$B$5,1,$B$8,$B$12)'; ws['A19']='Write p2'; ws['B19']='=XLL_CA_SET_INTERACTION_PARAMETERS_G(Start!$B$10,$B$4,$B$5,2,$B$9,$B$12)'; ws['A21']='Calculation token'; ws['B21']='=IF(AND(B18<>"",B19<>""),B12,B12)'; note(ws,23,'These SET formulas modify the working DAT. Keep an untouched baseline copy.',6,RED)
    ws=wb.create_sheet('Live targets'); sh(ws); head(ws,'Live thermodynamic target template','Slag-liq is forced as the entered phase. Blue target G values are illustrative and must be replaced by real targets for research use.'); hdr(ws,4,2,[('T, [C]',None,None),('P, [bar]',None,None),('IA',None,'Ca'),('IA',None,'Zn'),('IA',None,'O')]); states=[(1400,1,.8,.2,1),(1450,1,.6,.4,1),(1500,1,.4,.6,1),(1550,1,.2,.8,1),(1600,1,1,0,1)]
    for r,state in enumerate(states,8):
        for j,v in enumerate(state,2): ws.cell(r,j,v); inp(ws.cell(r,j))
    hdr(ws,4,8,[('G','Slag-liq',None),('ERROR',None,None)]); ws['M4']='Entered phases'; th(ws['M4']); ws['M5']='Slag-liq'; inp(ws['M5']); ws['H8']='=XLL_CA_CALCULATE(Start!$B$10,$B$4:$F$6,$B$8:$F$12,$H$4:$I$6,$M$5:$M$5,Parameters!$B$21)'; ws['J4']='Residual'; th(ws['J4']); ws['K4']='Target G (illustrative)'; th(ws['K4'])
    for r,t in zip(range(8,13),[-800000,-850000,-900000,-950000,-700000]): ws.cell(r,11,t); inp(ws.cell(r,11)); ws.cell(r,10,f'=IF(I{r}=0,K{r}-H{r},0)')
    note(ws,15,'The target numbers are illustrative; the sheet teaches the live derivative dependency, not a Ca-Zn-O assessment.',12,ORANGE)
    ws=wb.create_sheet('Live regression'); sh(ws); head(ws,'Use the live derivative matrix');
    for j,h in enumerate(['Weight','Residual','d/dp1','d/dp2'],1): ws.cell(4,j,h); th(ws.cell(4,j))
    for r in range(5,10): ws.cell(r,1,1); inp(ws.cell(r,1)); ws.cell(r,2,f"='Live targets'!J{r+3}"); ws.cell(r,3,f"='Derivative matrix setup'!D{r+3}"); ws.cell(r,4,f"='Derivative matrix setup'!E{r+3}")
    ws['A12']='Free mask'; ws['B12']=True; ws['C12']=True; ws['A13']='Fixed values'; ws['B13']=0; ws['C13']=0
    for c in ['B12','C12','B13','C13']: inp(ws[c])
    ws['A15']='Proposed Δp'; ws['B15']='=XLL_LINEAR_REGRESSION($C$5:$D$9,$B$5:$B$9,$A$5:$A$9,$B$12:$C$12,$B$13:$C$13)'; ws['A17']='Best 1 parameter'; ws['B17']='=XLL_BEST_COMBINATION($C$5:$D$9,$B$5:$B$9,$A$5:$A$9,$B$12:$C$12,$B$13:$C$13,1)'; ws['A20']='Suggested p1'; ws['B20']='=Parameters!B8+B15'; ws['A21']='Suggested p2'; ws['B21']='=Parameters!B9+C15'; note(ws,24,'Inspect and damp the proposed step before deliberately copying accepted values into Parameters. Do not create an automatic circular optimizer.',8)
    return save(wb,'03_PyroApp_Optimization_Workflows.xlsx')

def save(wb,name):
    # Native gridlines cover ordinary cells; thin borders keep filled headers,
    # input cells, and formulas visibly bounded as well.
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.__class__.__name__ != 'MergedCell' and cell.value is not None:
                    cell.border=CELL_BORDER
    p=OUT/name; wb.save(p); return p

files=[getting_started(),dat_parameters(),optimization()]
SUPPORTED_XLL_FUNCTIONS={
    'XLL_BEST_COMBINATION','XLL_CA_CALCULATE','XLL_CA_CHEMAPPDLLS','XLL_CA_DIMENSIONS',
    'XLL_CA_DIMENSIONS_MAX','XLL_CA_GET_COMPONENT_WEIGHTS','XLL_CA_GET_COMPOUND_CP',
    'XLL_CA_GET_COMPOUND_H298','XLL_CA_GET_COMPOUND_RANGE_COUNT','XLL_CA_GET_COMPOUND_S298',
    'XLL_CA_GET_COMPOUND_STOICHIOMETRY','XLL_CA_GET_COMPOUND_TUPPER','XLL_CA_GET_COMPOUND_WEIGHTS',
    'XLL_CA_GET_CONSTITUENT_CP','XLL_CA_GET_CONSTITUENT_H298','XLL_CA_GET_CONSTITUENT_RANGE_COUNT',
    'XLL_CA_GET_CONSTITUENT_S298','XLL_CA_GET_CONSTITUENT_STOICHIOMETRY',
    'XLL_CA_GET_CONSTITUENT_TUPPER','XLL_CA_GET_CONSTITUENT_WEIGHTS',
    'XLL_CA_GET_INTERACTION_INDICES','XLL_CA_GET_INTERACTION_PARAMETERS_G','XLL_CA_ISLITE',
    'XLL_CA_LICENSE_HOLDER_NAME','XLL_CA_LIST_COMPONENTS','XLL_CA_LIST_COMPOUNDS',
    'XLL_CA_LIST_CONSTITUENTS','XLL_CA_LIST_INTERACTIONS_G','XLL_CA_LIST_PHASES',
    'XLL_CA_LIST_SOLUTIONS','XLL_CA_LIST_SPECIES','XLL_CA_PROGRAM_ID','XLL_CA_SET_INTERACTION_PARAMETERS_G',
    'XLL_CA_USER_ID','XLL_CA_VERSION','XLL_DATA_CHANGE_BASIS','XLL_DATA_GENERATE_MESH',
    'XLL_LINEAR_REGRESSION','XLL_PD_ERROR_TABLE'
}
for p in files:
    wb=load_workbook(p,data_only=False); bad=[]; names=set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value=='=""': bad.append(f'{ws.title}!{c.coordinate}')
                if isinstance(c.value,str) and c.value.startswith('='): names.update(re.findall(r'\b(XLL_[A-Z0-9_]+)\s*\(',c.value))
    if bad: raise RuntimeError(f'{p.name}: legacy ="" markers at {bad}')
    unknown=names-SUPPORTED_XLL_FUNCTIONS
    if unknown: raise RuntimeError(f'{p.name}: unsupported XLL functions {sorted(unknown)}')
    print(p.name, len(wb.sheetnames),'sheets',len(names),'XLL functions',p.stat().st_size,'bytes')
